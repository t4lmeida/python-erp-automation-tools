"""
Extrator de Lançamentos Detalhados da Folha de Pagamento
=============================================================================
Dependências:
    pip install beautifulsoup4 openpyxl

Uso:
    python extrair_lancamentos_detalhados.py
"""

import re
import tkinter as tk
from tkinter import filedialog, ttk, messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from bs4 import BeautifulSoup
from typing import Optional

# ── Helpers ───────────────────────────────────────────────────────────────────

def carregar_elementos(caminho_html: str) -> list[dict]:
    with open(caminho_html, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    elementos = []
    for div in soup.find_all("div"):
        style = div.get("style", "")
        texto = div.get_text(strip=True).replace("\xa0", " ")
        if not texto:
            continue
        top_m  = re.search(r"top:\s*([\d.]+)px",  style)
        left_m = re.search(r"left:\s*([\d.]+)px", style)
        if top_m and left_m:
            t, l = float(top_m.group(1)), float(left_m.group(1))
            if t > 10:
                elementos.append({"top": t, "left": l, "text": texto})
    elementos.sort(key=lambda e: (e["top"], e["left"]))
    return elementos


def fv(lst):
    for e in lst:
        if re.search(r"\d", e["text"]):
            return e["text"]
    return ""


def converter_para_float(valor_str: str):
    """Remove pontos de milhar e substitui vírgula por ponto para converter em float."""
    if not valor_str:
        return None
    try:
        limpo = re.sub(r"[^\d.,]", "", str(valor_str).strip())
        limpo = limpo.replace(".", "").replace(",", ".")
        return float(limpo)
    except ValueError:
        return None


# ── Interface de Parâmetros Inicial ──────────────────────────────────────────

def obter_parametros_usuario() -> Optional[tuple[str, int, int]]:
    """Abre uma interface gráfica para coletar Empresa, Mês e Ano."""
    root = tk.Tk()
    root.title("Parâmetros do Extrator")
    root.geometry("320x240")
    root.resizable(False, False)

    # Centralizar janela na tela
    root.eval('tk::PlaceWindow . center')

    dados_informados = {}

    ttk.Label(root, text="Selecione a Empresa:", font=("Arial", 9, "bold")).pack(anchor="w", padx=20, pady=(15, 2))
    combo_empresa = ttk.Combobox(root, values=["Empresa1", "Empresa2"], state="readonly")
    combo_empresa.current(0)
    combo_empresa.pack(fill="x", padx=20)

    ttk.Label(root, text="Mês (1 a 12):", font=("Arial", 9, "bold")).pack(anchor="w", padx=20, pady=(10, 2))
    spin_mes = ttk.Spinbox(root, from_=1, to=12, width=10)
    spin_mes.set(1)
    spin_mes.pack(anchor="w", padx=20)

    ttk.Label(root, text="Ano (ex: 2026):", font=("Arial", 9, "bold")).pack(anchor="w", padx=20, pady=(10, 2))
    spin_ano = ttk.Spinbox(root, from_=2000, to=2100, width=10)
    spin_ano.set(2026)
    spin_ano.pack(anchor="w", padx=20)

    def confirmar():
        try:
            mes = int(spin_mes.get())
            ano = int(spin_ano.get())
            if not (1 <= mes <= 12):
                raise ValueError("Mês inválido")
            dados_informados["empresa"] = combo_empresa.get()
            dados_informados["mes"] = mes
            dados_informados["ano"] = ano
            root.destroy()
        except ValueError:
            messagebox.showerror("Erro de Validação", "Insira um mês e ano válidos.")

    ttk.Button(root, text="Avançar", command=confirmar).pack(pady=15)

    root.mainloop()

    if "empresa" in dados_informados:
        return dados_informados["empresa"], dados_informados["mes"], dados_informados["ano"]
    return None


# ── Extração ──────────────────────────────────────────────────────────────────

def extrair_lancamentos(elementos: list[dict], empresa: str, mes: int, ano: int) -> list[dict]:
    codigos = [
        e for e in elementos
        if 20 <= e["left"] <= 30 and re.match(r"^\d{5,}$", e["text"])
    ]

    lancamentos = []

    for idx, ce in enumerate(codigos):
        t = ce["top"]
        codigo = str(ce["text"]).strip()  # Garantindo o tipo string para a matrícula

        nome_els = [e for e in elementos if abs(e["top"] - t) <= 4 and 70 <= e["left"] <= 340]
        nome     = nome_els[0]["text"] if nome_els else ""

        sal_cont = fv([e for e in elementos if abs(e["top"] - t) <= 4 and 340 <= e["left"] <= 360])

        func_els = [e for e in elementos if abs(e["top"] - t) <= 4 and 460 <= e["left"] <= 610]
        funcao   = func_els[0]["text"] if func_els else ""

        admissao = fv([e for e in elementos if abs(e["top"] - (t + 17)) <= 5 and 455 <= e["left"] <= 475])

        tot_marker = [e for e in elementos if t + 50 <= e["top"] <= t + 200 and e["text"] == "*******"]
        t_tot = tot_marker[0]["top"] if tot_marker else t + 100

        lanc_tops = sorted(set(
            e["top"] for e in elementos
            if t + 30 <= e["top"] < t_tot - 2
            and 20 <= e["left"] <= 35
            and re.match(r"^\d{3}$", e["text"])
        ))

        for lt in lanc_tops:
            cod_lanc   = fv([e for e in elementos if abs(e["top"] - lt) <= 4 and 20 <= e["left"] <= 35])
            descricao  = " ".join(e["text"] for e in elementos if abs(e["top"] - lt) <= 4 and 50 <= e["left"] <= 285)
            referencia = fv([e for e in elementos if abs(e["top"] - lt) <= 5 and 285 <= e["left"] <= 305])
            provento   = fv([e for e in elementos if abs(e["top"] - lt) <= 4 and 415 <= e["left"] <= 430])
            desconto   = fv([e for e in elementos if abs(e["top"] - lt) <= 4 and 515 <= e["left"] <= 530])

            if provento:
                tipo, valor = "Provento", provento
            elif desconto:
                tipo, valor = "Desconto", desconto
            else:
                tipo, valor = "", ""

            lancamentos.append({
                "Empresa":         empresa,
                "Mês":             mes,
                "Ano":             ano,
                "Código Func.":    codigo,  # String
                "Nome":            nome,
                "Função":          funcao,
                "Admissão":        admissao,
                "Sal. Contratual": converter_para_float(sal_cont),
                "Cód. Lançamento": cod_lanc,
                "Descrição":       descricao,
                "Referência":      referencia,
                "Tipo":            tipo,
                "Valor":           converter_para_float(valor),
            })

    return lancamentos


# ── Geração do Excel ──────────────────────────────────────────────────────────

def gerar_excel(lancamentos: list[dict], caminho_xlsx: str) -> None:
    if not lancamentos:
        print("Nenhum lançamento encontrado.")
        return

    COLUNAS     = list(lancamentos[0].keys())
    MONEY_COLS  = {"Sal. Contratual", "Valor"}
    CENTER_COLS = {"Mês", "Ano", "Código Func.", "Admissão", "Cód. Lançamento", "Referência", "Tipo"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lançamentos Detalhados"

    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_font  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    prov_fill = PatternFill("solid", fgColor="E2EFDA")
    desc_fill = PatternFill("solid", fgColor="FCE4D6")
    alt_fill  = PatternFill("solid", fgColor="EEF4FB")
    brd = Border(left=Side(style="thin"), right=Side(style="thin"),
                 top=Side(style="thin"),  bottom=Side(style="thin"))

    # Cabeçalho
    for col, h in enumerate(COLUNAS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.border    = brd
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    prev_code     = None
    row_bg_toggle = False

    for i, lanc in enumerate(lancamentos):
        rn = i + 2
        if lanc["Código Func."] != prev_code:
            prev_code     = lanc["Código Func."]
            row_bg_toggle = not row_bg_toggle

        if lanc["Tipo"] == "Provento":
            fill = prov_fill
        elif lanc["Tipo"] == "Desconto":
            fill = desc_fill
        else:
            fill = alt_fill if row_bg_toggle else PatternFill("solid", fgColor="FFFFFF")

        for col, nome_col in enumerate(COLUNAS, 1):
            val  = lanc.get(nome_col, "")
            
            # Garantir escrita como texto explícito para a matrícula
            if nome_col == "Código Func.":
                cell = ws.cell(row=rn, column=col, value=str(val))
                cell.number_format = '@'
            else:
                cell = ws.cell(row=rn, column=col, value=val)

            cell.font   = Font(name="Arial", size=9)
            cell.fill   = fill
            cell.border = brd
            
            if nome_col in MONEY_COLS:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = 'R$ #,##0.00'
            elif nome_col in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    larguras = [15, 6, 6, 14, 30, 28, 11, 14, 12, 34, 10, 9, 12]
    for i, w in enumerate(larguras[:len(COLUNAS)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Legenda")
    ws2["A1"], ws2["B1"] = "Cor", "Significado"
    for c in ["A1", "B1"]:
        ws2[c].font = Font(bold=True)
    ws2["A2"].fill = prov_fill; ws2["B2"] = "Provento (crédito ao funcionário)"
    ws2["A3"].fill = desc_fill; ws2["B3"] = "Desconto (débito do funcionário)"
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 38

    wb.save(caminho_xlsx)
    
    print(f"✔ Arquivo salvo com sucesso: {caminho_xlsx}")
    print(f"  Empresa                          : {lancamentos[0]['Empresa']}")
    print(f"  Período                          : {lancamentos[0]['Mês']:02d}/{lancamentos[0]['Ano']}")
    print(f"  Total de lançamentos processados : {len(lancamentos)}")
    print(f"  Total de funcionários únicos     : {len(set(l['Código Func.'] for l in lancamentos))}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    # 1. Coletar Empresa, Mês e Ano do usuário
    params = obter_parametros_usuario()
    if not params:
        print("Operação cancelada pelo usuário.")
        return

    empresa, mes, ano = params

    # 2. Selecionar o arquivo HTML
    root = tk.Tk()
    root.withdraw()

    html_in = filedialog.askopenfilename(
        title="Selecione o arquivo HTML da folha de pagamento",
        filetypes=[("Arquivos HTML", "*.html *.htm"), ("Todos os arquivos", "*.*")]
    )
    if not html_in:
        print("Nenhum arquivo selecionado. Encerrando.")
        return

    print(f"Lendo e analisando estruturas: {html_in}")
    elementos   = carregar_elementos(html_in)
    lancamentos = extrair_lancamentos(elementos, empresa, mes, ano)

    if not lancamentos:
        print("Nenhum lançamento encontrado. Verifique o formato do HTML.")
        return

    # 3. Salvar o arquivo Excel
    xlsx_out = filedialog.asksaveasfilename(
        title="Salvar planilha de saída como...",
        defaultextension=".xlsx",
        initialfile=f"Lancamentos_{empresa}_{mes:02d}_{ano}.xlsx",
        filetypes=[("Planilha Excel", "*.xlsx"), ("Todos os arquivos", "*.*")]
    )
    if not xlsx_out:
        print("Nenhum destino selecionado. Encerrando.")
        return

    gerar_excel(lancamentos, xlsx_out)


if __name__ == "__main__":
    main()
